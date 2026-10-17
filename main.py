# -*- coding: utf-8 -*-
"""中国联通 · 日常任务合集 图形版 — 程序入口

用法:
    pythonw main.py                 图形界面 (推荐)
    python  main.py                 图形界面 (可看控制台错误)
    python  main.py --cli           命令行跑一次 (不开界面)
    python  main.py --selftest      自检: 配置/签名/加密/记录/导出
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

_BUS = None


def _bootstrap_log():
    """装日志重定向 (幂等); 必须先于任何会 print 的模块调用"""
    global _BUS
    if _BUS is not None:
        return _BUS
    from app import logstream, paths
    paths.ensure_dirs()
    _BUS = logstream.LogBus()
    logstream.install(_BUS)
    return _BUS


# ============================================================
# 自检 (不联网)
# ============================================================
def selftest():
    """不联网自检. 仅在 main() 已把 CMF_DATA_DIR 切到临时目录后调用."""
    import tempfile
    import threading

    from app import paths, records, scheduler
    from app.settings import MODULE_SWITCHES, AccountStore, Settings

    # 保险: 万一将来有人直接调用 selftest(), 再拦一道, 绝不在真实数据目录跑
    if "UnicomAllIn1_selftest" not in paths.data_dir():
        raise RuntimeError(
            "自检拒绝在真实数据目录运行: %s (请经由 main.py --selftest)" % paths.data_dir())

    ok, fail = [], []

    def chk(name, fn):
        try:
            fn()
            ok.append(name)
            print("[OK]   " + name)
        except Exception as e:
            import traceback
            traceback.print_exc()
            fail.append(name)
            print("[FAIL] %s -> %s" % (name, e))

    # 1 配置读写
    def t_settings():
        s = Settings()
        s.pushplus_token = "TEST_TOKEN"
        s.schedule_times = "06:30,19:45"
        s.enable_woread = True
        assert s.save()[0]
        s2 = Settings()
        assert s2.pushplus_token == "TEST_TOKEN"
        assert s2.schedule_times == "06:30,19:45"
        assert s2.enable_woread is True
        env = s2.as_env()
        assert env.get("PUSH_PLUS_TOKEN") == "TEST_TOKEN"
        assert env.get("UNICOM_TEST_MODE") == "0"
        s2.query_only = True
        assert s2.as_env().get("UNICOM_TEST_MODE") == "query"
        # 还原
        s2.pushplus_token = ""
        s2.schedule_times = "07:00,20:00"
        s2.enable_woread = False
        s2.query_only = False
        s2.save()
    chk("配置读写 / 环境变量转换", t_settings)

    # 2 globalConfig 转换
    def t_global_config():
        s = Settings()
        cfg = s.as_global_config()
        for key, _l, _h in MODULE_SWITCHES:
            assert key in cfg, "缺少开关 %s" % key
        assert isinstance(cfg["sign_config"], dict)
        assert isinstance(cfg["market_config"], dict)
        assert isinstance(cfg["regional_config"], dict)
        assert set(cfg["market_config"]) == {"run_task", "run_draw", "run_claim"}
        assert cfg["enable_notify"] in (True, False)
        assert cfg["refresh_device_id"] in (True, False)
        # 11 模块 + enable_notify + refresh_device_id + sign/market/regional 三个子配置
        assert len(cfg) == len(MODULE_SWITCHES) + 5, len(cfg)
    chk("globalConfig 结构转换 (11模块+子开关)", t_global_config)

    # 3 账号存取
    def t_accounts():
        st = AccountStore()
        st.accounts = []
        st.add("18600000000#pwd\naToken#appId&bToken")
        assert len(st.accounts) == 3, st.accounts
        assert st.add("18600000000#pwd") == 0, "重复账号应被忽略"
        assert len(st.enabled_configs()) == 3
        ids = [a.id for a in st.accounts]
        assert len(set(ids)) == 3, "账号 id 出现重复: %s" % ids
        st.toggle([st.accounts[0].id], False)
        assert len(st.enabled_configs()) == 2
        st.save()
        st2 = AccountStore()
        assert len(st2.accounts) == 3, "账号持久化失败"
        st.accounts, st2.accounts = [], []
        st.save()
        st2.save()
    chk("账号增删改查 / uuid 唯一性 / 持久化", t_accounts)

    # 4 日志总线
    def t_log():
        from app import logstream
        b = logstream.LogBus()
        w = logstream.StreamProxy(b)
        w.write("[03:15:01] 账号[1]测试A\n[03:15:02] 账号[2]测试B\n")
        w.flush()
        got = b.drain(10)
        assert len(got) == 2, got
        assert got[0]["account"] == 1 and got[1]["account"] == 2, \
            [(g["text"], g["account"]) for g in got]
        assert logstream.classify("❌ 失败") == "error"
        assert logstream.classify("🎉 抽奖成功") == "ok"
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False,
                                         encoding="utf-8") as f:
            p = f.name
        assert b.save_to(p) == 2
        os.unlink(p)
    chk("日志总线 (stdout重定向 / 账号归属 / 分级)", t_log)

    # 5 记录抽取: 资产
    def t_assets():
        cases = [
            ("[03:01:00] 账号[1]💰 [资产-话费] 当前余额: 23.50元, 实时话费: 1.00元",
             [("话费余额", "23.50"), ("实时话费", "1.00")]),
            ("[03:01:00] 账号[1]💰 [资产-阅读红包] 余额: 3.20元",
             [("阅读红包", "3.20")]),
            ("[03:01:00] 账号[1]💰 [资产-抢购] 本月权益超市话费累计: 12.00元",
             [("本月权益话费", "12.00")]),
            ("[03:01:00] 账号[1]云南生活: 💰 当前云豆余额: 88",
             [("云豆", "88")]),
            ("普通日志无资产", []),
        ]
        for text, expect in cases:
            assert records.extract_assets(text) == expect, (text, records.extract_assets(text))
    chk("资产字段抽取", t_assets)

    # 6 记录抽取: 中奖
    def t_wins():
        cases = [
            ("[03:01:00] 账号[1]权益超市: 🎉 抽奖成功: 5元话费", ("权益超市", "奖品", "5元话费")),
            ("[03:01:00] 账号[1]沃云手机: 抽奖获得：7天体验卡", ("沃云手机", "奖品", "7天体验卡")),
            ("[03:01:00] 账号[1]云南生活: 🎁 抽奖结果 - 云豆100", ("云南生活", "奖品", "云豆100")),
            ("[03:01:00] 账号[1]爱听任务: 获得 15", ("爱听任务", "积分", "15")),
            ("[03:01:00] 账号[1]🎉🎉🎉 [抢兑成功] 恭喜！已成功抢到目标奖品！ 🎉🎉🎉",
             ("抢兑成功", "奖品", None)),
            ("[03:01:00] 账号[1]权益超市: 💨 未中奖: xxx", None),
            ("[03:01:00] 账号[1]普通文本", None),
        ]
        for text, expect in cases:
            w = records.extract_win(text)
            if expect is None:
                assert w is None, (text, w)
            else:
                assert w is not None, text
                src, kind = expect[0], expect[1]
                assert (w["来源"], w["类型"]) == (src, kind), (text, w)
                if expect[2] is not None:
                    assert w["内容"] == expect[2], (text, w)
    chk("中奖记录抽取 (10 条规则)", t_wins)

    # 7 记录去重 + Excel 导出
    def t_records():
        records.clear()
        row = {"时间": "12:01:02", "账号": 1, "手机号": "186****0000",
               "来源": "权益超市", "类型": "奖品", "内容": "5元话费", "账号ID": "acc1"}
        assert records.ingest_wins([dict(row)]) == 1
        assert records.ingest_wins([dict(row)]) == 0, "重复记录应被去重"
        assert records.ingest_wins([dict(row, 内容="10元话费")]) == 1
        assert records.count() == 2
        res = [{"时间": "2026-10-05 08:00", "账号": 1, "手机号": "186****0000",
                "状态": "完成", "话费余额": "¥23.50", "本次收益": "5元话费"}]
        records.append_results(res)
        assert len(records.all_results()) == 1

        out = os.path.join(tempfile.gettempdir(), "uni_wins_test.xlsx")
        n = records.export_wins(out, records.all_wins())
        assert n == 2 and os.path.getsize(out) > 3000
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb.active
        assert ws.max_row == 3 and ws.max_column == 6, (ws.max_row, ws.max_column)
        assert ws.cell(1, 1).value == "时间"
        os.unlink(out)

        out2 = os.path.join(tempfile.gettempdir(), "uni_res_test.xlsx")
        records.export_results(out2, records.all_results())
        assert os.path.getsize(out2) > 2500
        os.unlink(out2)
        records.clear()
    chk("记录去重 / 运行结果 / Excel 导出", t_records)

    # 8 定时解析
    def t_sched():
        assert scheduler.parse_times("07:00, 20:00") == ["07:00", "20:00"]
        assert scheduler.parse_times("7:5,25:00,x,07:00") == ["07:00", "07:05"]
        assert scheduler.parse_times("") == []
        assert scheduler.parse_times("7:5,7:05,07:05") == ["07:05"]
        assert scheduler.parse_times("24:00,99:99") == []
    chk("定时时刻解析 / 归一化去重", t_sched)

    # 9 核心模块导入 + 常量
    def t_core():
        from core import legacy
        assert legacy.SCRIPT_VERSION == "v1.2.1"
        # v1.1.2 的 YPHD_SECRET_KEY 在 v1.2.1 已重命名为 HOMETOWN_LOTTERY_SECRET
        assert legacy.HOMETOWN_LOTTERY_SECRET == "s8Hf3LqP9xN2vM5bR7tY1wZ4cA6eG0K"
        assert callable(legacy.UserService)
        for m in ("execute_daily_tasks", "ensure_login", "sign_task_main",
                  "market_task", "ltyp_task", "queryRemain", "do_notify"):
            assert hasattr(legacy.UserService, m) or hasattr(legacy, m), m
        assert callable(legacy.cross_view_security_share_keys)
        assert callable(legacy._check_stop)
        # v1.2.1 新增: 海岛逐浪 / AI入戏
        assert hasattr(legacy, "ISLAND_LOCAL_IMAGES")
        assert hasattr(legacy, "ISLAND_PRIZE_CACHE_PATH")
        assert legacy.ISLAND_ACTIVITY_ID == "NDA="
    chk("核心模块导入 / 版本与签名常量", t_core)

    # 10 加解密 (pycryptodome)
    def t_crypto():
        from Crypto.Cipher import AES, PKCS1_v1_5
        from Crypto.PublicKey import RSA
        from Crypto.Util.Padding import pad, unpad
        key, iv = b"0123456789abcdef", b"wNSOYIB1k1DjY5lA"
        plain = '{"clientId":"1001000003"}'
        c = AES.new(key, AES.MODE_CBC, iv)
        enc = c.encrypt(pad(plain.encode(), 16))
        assert unpad(AES.new(key, AES.MODE_CBC, iv).decrypt(enc), 16).decode() == plain
        # 核心脚本用的 RSA 登录公钥必须可解析
        from core import legacy
        k = RSA.import_key(legacy.LOGIN_PUB_KEY)
        assert k.size_in_bits() == 1024
        cipher = PKCS1_v1_5.new(k)
        assert len(cipher.encrypt(b"12345678")) == 128
    chk("AES / RSA 加解密 (登录链依赖)", t_crypto)

    # 11 停止控制
    def t_stop():
        from core import legacy
        ev = threading.Event()
        legacy.STOP_EVENT = ev
        assert legacy._stopped() is False
        assert legacy._check_stop(" (测试)") is False
        ev.set()
        assert legacy._stopped() is True
        assert legacy._check_stop(" (测试)") is True
        assert legacy._sleep(5) is True, "停止信号应立即打断 sleep"
        legacy.STOP_EVENT = None
        assert legacy._stopped() is False
    chk("停止控制 (含可中断等待)", t_stop)

    # 12 汇总 HTML + notify
    def t_summary():
        from app import runner
        html = runner.build_summary_html(
            [{"序号": 1, "手机号": "186****0000", "状态": "完成",
              "话费余额": "¥23.50", "本次收益": "5元话费"}], 3)
        assert "日常任务汇总" in html and "</table>" in html
        assert "话费余额" in html and "5元话费" in html
        import notify
        assert callable(notify.send)
    chk("汇总表格生成 / notify 模块", t_summary)

    # 13 人脸图清单 → 环境变量 (v1.2.1 海岛逐浪/AI入戏)
    def t_face_images():
        from app.settings import Settings
        s = Settings()
        s.face_images = []
        assert s.face_images == []
        # 清空后不应再导出该键 (避免旧值残留)
        assert "UNICOM_CMF_LOCAL_IMAGES" not in s.as_env()
        s.face_images = r"C:\a\1.jpg"
        assert s.face_images == [r"C:\a\1.jpg"]
        assert s.as_env()["UNICOM_CMF_LOCAL_IMAGES"] == r"C:\a\1.jpg"
        # 逗号/中文逗号分隔的字符串应被拆成列表
        s.face_images = r"C:\a\1.jpg,D:\b\2.png"
        assert s.face_images == [r"C:\a\1.jpg", r"D:\b\2.png"], s.face_images
        assert s.as_env()["UNICOM_CMF_LOCAL_IMAGES"] == r"C:\a\1.jpg,D:\b\2.png"
        s.face_images = "C:\\a\\1.jpg，D:\\b\\2.png"
        assert s.face_images == [r"C:\a\1.jpg", r"D:\b\2.png"], s.face_images
        s.face_images = []
        assert s.as_env().get("UNICOM_CMF_DEL_CLOUD_PHOTOS") == "1"
        s.cmf_del_cloud_photos = False
        assert s.as_env()["UNICOM_CMF_DEL_CLOUD_PHOTOS"] == "0"
        s.cmf_del_cloud_photos = True
        # 持久化后应保持为列表
        s.face_images = [r"C:\a\1.jpg"]
        assert s.save()[0]
        s2 = Settings()
        assert isinstance(s2.face_images, list), type(s2.face_images)
        assert s2.face_images == [r"C:\a\1.jpg"], s2.face_images
        s2.face_images = []
        s2.save()
    chk("人脸图清单 列表/环境变量/持久化", t_face_images)

    print()
    print("=" * 50)
    print("自检完成: 通过 %d / 失败 %d" % (len(ok), len(fail)))
    if fail:
        print("失败项: " + ", ".join(fail))
    print("=" * 50)
    return 0 if not fail else 1


# ============================================================
# 命令行模式
# ============================================================
def cli_run():
    import threading

    from app import paths, runner
    from app.settings import AccountStore, Settings

    paths.ensure_dirs()
    settings, store = Settings(), AccountStore()
    if not store.enabled_configs():
        print("❌ 没有配置账号, 请先运行图形界面添加账号")
        return 1

    class ConsoleBus(object):
        def emit(self, text, level=None, account=None):
            print(text)

        def drain(self, limit=0):
            return []

        def snapshot(self):
            return []

    res = runner.run_batch(settings, store, ConsoleBus(), threading.Event())
    print("\n完成: %d 账号, 新增中奖记录 %d 条, 推送=%s"
          % (len(res.get("results", [])), res.get("wins", 0), res.get("pushed")))
    return 0


def main():
    args = [a.lower() for a in sys.argv[1:]]

    # ⚠️ 自检会增删账号/配置/中奖记录, 必须在导入任何 app.* 之前
    #    把数据目录切到临时目录 —— paths 的常量在模块导入时就解析完成,
    #    必须先于 _bootstrap_log() (它会 import app.paths)。
    if "--selftest" in args or "--self-test" in args:
        import tempfile
        os.environ["CMF_DATA_DIR"] = os.path.join(
            tempfile.gettempdir(), "UnicomAllIn1_selftest")
        if sys.stdout is None:
            _bootstrap_log()
        sys.exit(selftest())

    # 打包成 windowed exe 时 sys.stdout 为 None, print 会崩 → 先装日志重定向
    if sys.stdout is None:
        _bootstrap_log()

    if "--cli" in args:
        sys.exit(cli_run())

    _bootstrap_log()
    try:
        from app import gui
        gui.run()
    except SystemExit:
        raise
    except BaseException:
        import traceback
        err = sys.__stderr__ or sys.__stdout__
        try:
            traceback.print_exc(file=err)
            err.write("\n程序异常退出。按回车关闭...\n")
            err.flush()
            input()
        except Exception:
            pass
        sys.exit(1)


if __name__ == "__main__":
    main()
