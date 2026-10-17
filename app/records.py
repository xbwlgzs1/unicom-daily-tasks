# -*- coding: utf-8 -*-
"""记录库: 从运行日志里抽取【资产】与【中奖/获取】, 本地持久化并可导出 Excel.

原脚本只把结果 print 出来, 不落盘; 这里用正则按行解析它的输出格式,
因此无需改动那 7000 行核心代码, 也能拿到结构化数据.
"""
import json
import os
import re
import threading
import time
from datetime import datetime

from . import paths

_lock = threading.Lock()
MAX_RECORDS = 5000

# ------------------------------------------------------------
# 资产: 单账号最新值 (key -> value)
# ------------------------------------------------------------
ASSET_RULES = [
    ("话费余额",   re.compile(r"\[资产-话费\]\s*当前余额:\s*([-\d.]+)元")),
    ("实时话费",   re.compile(r"实时话费:\s*([-\d.]+)元")),
    ("阅读红包",   re.compile(r"\[资产-阅读红包\]\s*余额:\s*([-\d.]+)元")),
    ("本月权益话费", re.compile(r"\[资产-抢购\]\s*本月权益超市话费累计:\s*([-\d.]+)元")),
    ("云豆",       re.compile(r"当前云豆余额:\s*(\S+)")),
]

# ------------------------------------------------------------
# 中奖/获取: 历史记录
# (来源, 正则, 取值组; None = 整行作为内容)
# ------------------------------------------------------------
WIN_RULES = [
    ("抢兑成功",   re.compile(r"\[抢兑成功\]"), None, "奖品"),
    ("抢兑明细",   re.compile(r"🎁\s*\[抢兑\]\s*(\S+?)\s*\|\s*(.+?)\s*\(变动:\s*(.+?)\)"), None, "奖品"),
    ("权益超市",   re.compile(r"权益超市:\s*🎉\s*抽奖成功[:：]\s*(.+)"), 1, "奖品"),
    ("沃云手机",   re.compile(r"沃云手机:\s*抽奖获得[:：]\s*(.+)"), 1, "奖品"),
    ("云南生活",   re.compile(r"云南生活:\s*🎁\s*抽奖结果\s*-\s*(.+)"), 1, "奖品"),
    ("河南商都",   re.compile(r"河南商都:.*?获得\s+(.+)"), 1, "奖品"),
    ("安徽超级星期五", re.compile(r"安徽超级星期五:\s*🎉\s*第\d+次抢购成功"), None, "奖品"),
    ("权益超市-会员中心", re.compile(r"权益超市-会员中心:\s*✅\s*(.+)"), 1, "积分"),
    ("爱听任务",   re.compile(r"爱听任务:\s*获得\s+(.+)"), 1, "积分"),
    ("联通爱听",   re.compile(r"联通爱听:.*?获得\s+(.+)"), 1, "积分"),
]

WIN_COLUMNS = ["时间", "账号", "手机号", "来源", "类型", "内容"]
RESULT_COLUMNS = ["时间", "账号", "手机号", "状态", "话费余额", "本次收益"]


def _load(path, default=None):
    if not os.path.exists(path):
        return default if default is not None else []
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else (default if default is not None else [])
    except Exception:
        return default if default is not None else []


def _save(path, data):
    paths.ensure_dirs()
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ------------------------------------------------------------
# 抽取
# ------------------------------------------------------------
def extract_assets(text):
    """一行日志里可能带资产信息 -> [(名称, 数值)]"""
    out = []
    for name, rx in ASSET_RULES:
        m = rx.search(text)
        if m:
            out.append((name, m.group(1)))
    return out


def extract_win(text):
    """命中中奖/获取规则 -> dict, 未命中返回 None"""
    for source, rx, grp, kind in WIN_RULES:
        m = rx.search(text)
        if not m:
            continue
        if source == "抢兑明细":
            # 三个捕获组: 时间 | 内容 (变动:金额)
            content = "%s (变动:%s)" % (m.group(2), m.group(3))
            timestr = m.group(1)
        elif grp is None:
            content = re.sub(r"[\U0001F389\U0001F3C6\s]+", "", text).strip() or text.strip()
            timestr = ""
        else:
            content = m.group(grp).strip()
            timestr = ""
        if not content:
            continue
        return {"来源": source, "类型": kind, "内容": content[:200], "时间": timestr}
    return None


# ------------------------------------------------------------
# 写入
# ------------------------------------------------------------
def ingest_wins(new_records):
    """去重写入; 返回新增条数. key = 账号|来源|内容|时间(分钟)"""
    if not new_records:
        return 0
    with _lock:
        old = _load(paths.WINS_FILE)
        seen = {r.get("_k") for r in old if isinstance(r, dict)}
        added = 0
        for r in new_records:
            key = "|".join([str(r.get("账号ID", "")), str(r.get("来源", "")),
                            str(r.get("内容", "")), str(r.get("时间", ""))[:16]])
            if key in seen:
                continue
            seen.add(key)
            r["_k"] = key
            old.append(r)
            added += 1
        if added:
            old.sort(key=lambda x: x.get("时间") or "", reverse=True)
            del old[MAX_RECORDS:]
            _save(paths.WINS_FILE, old)
    return added


def all_wins():
    with _lock:
        data = _load(paths.WINS_FILE)
    data.sort(key=lambda x: x.get("时间") or "", reverse=True)
    return data


def append_results(rows):
    """每次运行追加一行运行结果 (保留最近 500 条)"""
    if not rows:
        return 0
    with _lock:
        old = _load(paths.RESULTS_FILE)
        old = [r for r in old if isinstance(r, dict)]
        old.extend(rows)
        old.sort(key=lambda x: x.get("时间") or "", reverse=True)
        del old[500:]
        _save(paths.RESULTS_FILE, old)
    return len(rows)


def all_results():
    with _lock:
        data = _load(paths.RESULTS_FILE)
    data.sort(key=lambda x: x.get("时间") or "", reverse=True)
    return data


def clear(kind="all"):
    with _lock:
        if kind in ("all", "wins"):
            _save(paths.WINS_FILE, [])
        if kind in ("all", "results"):
            _save(paths.RESULTS_FILE, [])


def count():
    with _lock:
        return len(_load(paths.WINS_FILE))


# ------------------------------------------------------------
# Excel 导出
# ------------------------------------------------------------
def _export(path, records, columns, sheet):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = sheet

    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(columns)
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for r in records:
        ws.append([str(r.get(c, "") or "") for c in columns])

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="center")

    widths = [18, 10, 15, 18, 8, 60][:len(columns)]
    while len(widths) < len(columns):
        widths.append(16)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(records)


def export_wins(path, records=None):
    records = all_wins() if records is None else records
    if path.lower().endswith(".csv"):
        return _export_csv(path, records, WIN_COLUMNS)
    return _export(path, records, WIN_COLUMNS, "中奖记录")


def export_results(path, records=None):
    records = all_results() if records is None else records
    if path.lower().endswith(".csv"):
        return _export_csv(path, records, RESULT_COLUMNS)
    return _export(path, records, RESULT_COLUMNS, "运行结果")


def _export_csv(path, records, columns):
    import csv
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in records:
            w.writerow({c: str(r.get(c, "") or "") for c in columns})
    return len(records)
